# -*- coding: utf-8 -*-
"""The example files, and the school they describe.

The product asks for "a clear example of a correctly formatted input file".
The example here is not a picture of one: it is a working school, small
enough to read in a minute and complete enough to solve, and it is produced
by the same code the download button serves.  So an administrator can
download it, upload it straight back, and watch a timetable come out — which
is the only demonstration of a file format anyone believes.

`spec_to_csv` writes any school in the same shape, so a school that was
imported once can be exported, edited in Excel and re-imported.
"""
from __future__ import annotations

import csv
import io

from app import importer, spec as spec_mod

#: Excel opens a UTF-8 CSV as mojibake unless it is given a byte-order mark,
#: and Hebrew makes that immediately obvious, so every file written here has
#: one.  CRLF for the same reason.
_BOM = "﻿"


def _csv(blocks: list[tuple[str, list[list]]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    for n, (title, rows) in enumerate(blocks):
        if n:
            writer.writerow([])
        writer.writerow([f"## {title}"])
        for row in rows:
            writer.writerow(["" if c is None else c for c in row])
    return (_BOM + buf.getvalue()).encode("utf-8")


# ---------------------------------------------------------- the example

_TEACHERS = [
    ["שם המורה", "שעות שבועיות", "מחנכת של", "מקסימום שעות ליום",
     "מינימום שעות ליום", "ימי עבודה", "מקסימום חלונות"],
    ["רות כהן", 15, "א", 5, 3, "", 1],
    ["שרה לוי", 15, "ב", 5, 3, "", 1],
    ["דנה מזרחי", 12, "", 4, 3, "", 2],
    ["מיכל אבן", 8, "", 4, 3, 2, 2],
]

_CLASSES = [
    ["שם הכיתה", "יום", "מינימום שעות", "מקסימום שעות", "שעות מדויקות"],
    ["א", "כל הימים", 5, 5, 5],
    ["ב", "כל הימים", 5, 5, 5],
]

_LESSONS = [
    ["כיתה", "מקצוע", "מורות", "שעות", "פיצול"],
    ["א", "תורה", "רות כהן", 5, ""],
    ["א", "חשבון", "רות כהן", 5, "2,1,1,1"],
    ["א", "עברית", "רות כהן", 4, ""],
    ["א", "חברה", "רות כהן", 1, ""],
    ["א", "אנגלית", "דנה מזרחי", 3, ""],
    ["א", "מדעים", "דנה מזרחי", 3, ""],
    ["א", "חינוך גופני", "מיכל אבן", 2, ""],
    ["א", "אומנות", "מיכל אבן", 2, ""],
    ["ב", "תורה", "שרה לוי", 5, ""],
    ["ב", "חשבון", "שרה לוי", 5, "2,1,1,1"],
    ["ב", "עברית", "שרה לוי", 4, ""],
    ["ב", "חברה", "שרה לוי", 1, ""],
    ["ב", "אנגלית", "דנה מזרחי", 3, ""],
    ["ב", "מדעים", "דנה מזרחי", 3, ""],
    ["ב", "חינוך גופני", "מיכל אבן", 2, ""],
    ["ב", "אומנות", "מיכל אבן", 2, ""],
]

_DAYS = [
    ["יום", "מספר שעות"],
    ["ראשון", 6],
    ["שני", 6],
    ["שלישי", 6],
    ["רביעי", 6],
    ["חמישי", 6],
]

_LIMITS = [
    ["שם המורה", "סוג", "יום", "שעות", "הערה"],
    ["דנה מזרחי", "יום חופש", "שני", "", "אינה מלמדת בשני"],
    ["מיכל אבן", "בחירת יום חופש", "שלישי, רביעי", "",
     "יום חופש באחד משני הימים"],
    ["רות כהן", "שעות חסומות", "כל הימים", 6, "אינה מלמדת בשעה השישית"],
    ["שרה לוי", "סיום מוקדם", "חמישי", 4, "מסיימת בשעה 4 בחמישי"],
]

#: The two files the product asks the administrator for.  The first carries
#: the staff, the classes and what is taught; the second, when people cannot
#: teach.  Both are the same format, so either may carry either table.
TEACHERS_FILE = "1-מורות-כיתות-ושיעורים.csv"
LIMITS_FILE = "2-ימים-ושעות-חסומים.csv"


WORD_FILE = "בית-ספר-לדוגמה.docx"

#: The four tables of the first file, in the order they are taught.
_BLOCKS = [("ימים", _DAYS), ("מורות", _TEACHERS), ("כיתות", _CLASSES),
           ("שיעורים", _LESSONS)]


def template_files() -> list[tuple[str, bytes]]:
    return [
        (TEACHERS_FILE, _csv(_BLOCKS)),
        (LIMITS_FILE, _csv([("אילוצים", _LIMITS)])),
    ]


def template_docx() -> bytes | None:
    """The same example as one Word document — a heading per table.

    Returns None where python-docx is not installed.  Writing .docx is the
    one thing in this program that needs it; *reading* one does not, so an
    office without the package can still upload its own Word file and simply
    does not get this particular example in that format.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
    except ImportError:
        return None
    import io as _io

    doc = Document()
    doc.styles["Normal"].font.size = __import__(
        "docx.shared", fromlist=["Pt"]).Pt(9)

    def rtl(par):
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        par._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
        return par

    rtl(doc.add_heading("קובץ נתונים לדוגמה — מערכת שעות", level=0))
    rtl(doc.add_paragraph(
        "כותרת לכל טבלה, והטבלה מתחתיה. שמות העמודות הם מה שקובע — הסדר "
        "אינו חשוב, ועמודה שאינה רלוונטית אפשר להשמיט. אפשר להעלות את הקובץ "
        "הזה כפי שהוא."))
    for title, rows in _BLOCKS + [("אילוצים", _LIMITS)]:
        rtl(doc.add_heading(title, level=1))
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        for n, row in enumerate(rows):
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = "" if value is None else str(value)
                for par in cell.paragraphs:
                    rtl(par)
                    for run in par.runs:
                        run.bold = n == 0
    buf = _io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def demo_spec() -> tuple[spec_mod.SchoolSpec | None, importer.ImportReport]:
    """The example files, read back through the real importer."""
    return importer.read_files(template_files(),
                               name="בית ספר לדוגמה", year="תשפ\"ז")


#: What each column means, for the screen that shows the format.  Kept beside
#: the example so the two cannot drift apart.
COLUMN_HELP = [
    ("איך מסמנים טבלה",
     "ב-CSV — שורה שמתחילה ב-##, למשל '## מורות'. ב-Word — כותרת רגילה "
     "מעל הטבלה. ב-Excel — שם הגיליון. אם אין סימון, המערכת מזהה את "
     "הטבלה לפי שמות העמודות שלה."),
    ("ימים", "מבנה השבוע: שם כל יום ומספר השעות בו. אפשר להשמיט — "
             "ברירת המחדל היא ראשון–שישי עם שישי מקוצר."),
    ("מורות", "שורה לכל מורה. חובה: שם המורה. השאר אופציונלי — "
              "'מחנכת של' קושר מורה לכיתה ומפעיל את כלל פתיחת היום. "
              "מורה מקצועית המלמדת כמה כיתות זקוקה לרוב ל'מקסימום חלונות' 2: "
              "כשכל כיתה לומדת ברצף, מי שנכנסת לשתי כיתות באותו יום נדחקת "
              "לאמצע היום של אחת מהן."),
    ("כיתות", "אורך יום הלימודים לכל כיתה. 'יום' יכול להיות 'כל הימים' "
              "או יום מסוים, כדי לקצר יום אחד."),
    ("שיעורים", "מי מלמדת מה ולכמה שעות בשבוע. שתי מורות המלמדות במקביל "
                "(הקבצה) מופרדות ב-'|'. 'פיצול' קובע שיעורים כפולים: "
                "'2,1,1,1' הוא כפול ועוד שלושה בודדים."),
    ("אילוצים", "מתי מורה אינה זמינה. סוגים: יום חופש · בחירת יום חופש · "
                "שעות חסומות · סיום מוקדם."),
]


# ------------------------------------------------------------ export back

def spec_to_csv(spec: spec_mod.SchoolSpec) -> list[tuple[str, bytes]]:
    """Write a school back out in the format it can be read from."""
    days = [["יום", "מספר שעות"]] + [
        [name, spec.grid.periods_per_day[d]]
        for d, name in enumerate(spec.grid.days)]

    teachers = [_TEACHERS[0]]
    for t in spec.teachers:
        teachers.append([t.name, t.declared or "", t.homeroom or "",
                         t.max_per_day or "", t.min_per_day,
                         t.exact_working_days or "", t.max_windows])

    classes = [_CLASSES[0]]
    for c in spec.classes:
        # One row per distinct day shape, so an unchanged week stays one line.
        shapes: dict[tuple, list[int]] = {}
        for d in range(len(spec.grid.days)):
            lo, hi = c.day_span.get(d, (0, spec.grid.periods_per_day[d]))
            shapes.setdefault((lo, hi, c.preferred.get(d)), []).append(d)
        for (lo, hi, exact), ds in shapes.items():
            label = ("כל הימים" if len(ds) == len(spec.grid.days)
                     else ", ".join(spec.grid.days[d] for d in ds))
            classes.append([c.name, label, lo, hi, exact if exact else ""])

    lessons = [_LESSONS[0]]
    for lesson in spec.lessons:
        lessons.append([lesson.klass or "", lesson.subject,
                        "|".join(lesson.teachers), lesson.hours,
                        ",".join(str(x) for x in lesson.pattern)])

    limits = [_LIMITS[0]]
    for t in spec.teachers:
        for d in t.off_fixed:
            limits.append([t.name, "יום חופש", spec.grid.days[d], "", ""])
        if t.off_choice:
            limits.append([t.name, "בחירת יום חופש",
                           ", ".join(spec.grid.days[d] for d in t.off_choice),
                           "", ""])
        if t.forbidden_periods:
            limits.append([t.name, "שעות חסומות", "כל הימים",
                           ",".join(str(p) for p in t.forbidden_periods), ""])
        for d, last in sorted(t.latest_period_on.items()):
            limits.append([t.name, "סיום מוקדם", spec.grid.days[int(d)],
                           last, ""])

    # The exported file names the school; the blank template does not, so
    # the format a first-time reader sees stays as short as possible.
    head = [["שם בית הספר", "שנה"], [spec.name, spec.year]]
    return [
        (TEACHERS_FILE, _csv([("בית ספר", head), ("ימים", days),
                              ("מורות", teachers), ("כיתות", classes),
                              ("שיעורים", lessons)])),
        (LIMITS_FILE, _csv([("אילוצים", limits)])),
    ]
