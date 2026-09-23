# -*- coding: utf-8 -*-
"""Tests for the platform layer.

The engine's own tests (`test_scheduler.py`) prove the rules are modelled
right.  These prove the layer around them: that a school survives a trip
through a file, that installing one into the engine leaves nothing of the
last one behind, that the fast evaluator agrees with the solver, and that the
editing tools never hand the user a placement that breaks a hard rule.

Run with: python -m unittest discover -s tests -t .
"""
from __future__ import annotations

import json
import pathlib
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

import checks
import model
import school
import solver
from model import Timetable, expand
from app import activate
from app import alternatives as alt
from app import constraints as cat
from app import engine
from app import evaluate
from app import importer
from app import samples
from app import spec as spec_mod
from app import store as store_mod
from app import timetable as tt_mod

ROOT = pathlib.Path(__file__).resolve().parents[1]


def _demo():
    spec, report = samples.demo_spec()
    assert spec is not None, report.errors
    return spec


class TestGrid(unittest.TestCase):
    """The week is derived by counting, not named."""

    def test_the_real_school_short_day_is_friday_and_nothing_else(self):
        self.assertEqual(model.full_days(), [0, 1, 2, 3, 4])
        self.assertEqual(model.short_days(), [model.FRI])

    def test_a_week_with_no_short_day_has_no_short_days(self):
        spec = _demo()
        with activate.activated(spec):
            self.assertEqual(model.short_days(), [])
            self.assertEqual(len(model.full_days()), 5)

    def test_the_grid_is_restored_afterwards(self):
        before = list(model.DAYS), list(model.PERIODS_PER_DAY)
        with activate.activated(_demo()):
            pass
        self.assertEqual((list(model.DAYS), list(model.PERIODS_PER_DAY)),
                         before)


class TestActivation(unittest.TestCase):
    def test_every_school_constant_is_installed_and_restored(self):
        """A new constant in `school.py` that nobody adds to the activator
        would leak from one school into the next, silently."""
        exempt = {
            # Day indices imported from `model`; not school data.
            "SUN", "MON", "TUE", "WED", "THU", "FRI",
            # Derived inside `school.py` and read nowhere else.
            "CLASS_NAMES",
        }
        public = {n for n in dir(school)
                  if n.isupper() and not n.startswith("_")} - exempt
        self.assertEqual(public - set(activate._SCHOOL_ATTRS), set())

    def test_globals_come_back(self):
        before = {a: getattr(school, a) for a in activate._SCHOOL_ATTRS}
        with activate.activated(_demo()):
            self.assertEqual([t.name for t in school.TEACHERS],
                             ["רות כהן", "שרה לוי", "דנה מזרחי", "מיכל אבן"])
        for attr, value in before.items():
            self.assertIs(getattr(school, attr), value, attr)

    def test_nesting_restores_each_level(self):
        outer, inner = _demo(), _demo()
        inner.name = "פנימי"
        with activate.activated(outer):
            names = list(model.DAYS)
            with activate.activated(inner):
                pass
            self.assertEqual(list(model.DAYS), names)


class TestSpecRoundTrip(unittest.TestCase):
    """The built-in school, captured as data, has to mean the same thing."""

    @classmethod
    def setUpClass(cls):
        cls.captured = activate.capture()
        cls.reloaded = spec_mod.from_dict(json.loads(json.dumps(
            spec_mod.to_dict(cls.captured), ensure_ascii=False)))

    def test_it_is_well_formed(self):
        self.assertEqual(self.reloaded.problems(), [])

    def test_the_requirements_are_identical(self):
        want = {(r.klass, r.subject, r.teachers, r.hours, r.pattern)
                for r in school.requirements()}
        with activate.activated(self.reloaded):
            got = {(r.klass, r.subject, r.teachers, r.hours, r.pattern)
                   for r in self.reloaded.requirements()}
        self.assertEqual(got, want)
        self.assertEqual(len(want), 187)

    def test_phase_zero_says_the_same_thing(self):
        want = [str(f) for f in checks.run(school.requirements())]
        with activate.activated(self.reloaded):
            got = [str(f) for f in checks.run(self.reloaded.requirements())]
        self.assertEqual(got, want)


class TestEvaluatorAgreesWithTheSolver(unittest.TestCase):
    """The fast Python pricing is a second implementation of the priced
    rules.  It is only worth having if it agrees with the one that counts."""

    @classmethod
    def setUpClass(cls):
        cls.spec = activate.capture()
        saved = json.loads((ROOT / "solution.json").read_text(encoding="utf-8"))
        cls.placement = {int(k): tuple(v)
                         for k, v in saved["placement"].items()}

    def test_the_total_and_every_label_match(self):
        with activate.activated(self.spec):
            reqs = self.spec.requirements()
            meetings = expand(reqs)
            mine = evaluate.evaluate(self.spec, cat.Settings(), self.placement,
                                     meetings, reqs)
            theirs = solver.score(reqs, self.placement, time_limit=120,
                                  workers=8)
        self.assertIsNotNone(theirs, "the approved timetable must be legal")
        self.assertEqual(mine.total, sum(theirs.values()))

        # The fifth-day components are split between the labels by *which*
        # day carries the rule, and where two days cost the same the solver
        # is free to pick either.  So those labels are compared as one sum
        # and the rest one by one.
        fifth = lambda d: sum(v for k, v in d.items()
                              if k.startswith("יום חמישי:"))
        rest = lambda d: {k: v for k, v in d.items()
                          if not k.startswith("יום חמישי:")}
        self.assertEqual(fifth(mine.penalties), fifth(theirs))
        self.assertEqual(rest(mine.penalties), rest(theirs))

    def test_the_approved_timetable_breaks_no_hard_rule(self):
        with activate.activated(self.spec):
            meetings = expand(self.spec.requirements())
            got = evaluate.evaluate(self.spec, cat.Settings(), self.placement,
                                    meetings)
        self.assertEqual(got.violations, [])
        self.assertEqual(got.unplaced, [])


class TestImporter(unittest.TestCase):
    def test_the_example_files_describe_a_usable_school(self):
        spec, report = samples.demo_spec()
        self.assertTrue(report.ok, report.errors)
        self.assertEqual(report.warnings, [])
        self.assertEqual(spec.problems(), [])
        self.assertEqual(len(spec.teachers), 4)
        self.assertEqual(sum(l.hours for l in spec.lessons), 50)

    def test_every_kind_of_limit_lands_where_it_belongs(self):
        spec = _demo()
        self.assertEqual(spec.teacher("דנה מזרחי").off_fixed, [1])
        self.assertEqual(spec.teacher("מיכל אבן").off_choice, [2, 3])
        self.assertEqual(spec.teacher("רות כהן").forbidden_periods, [6])
        self.assertEqual(spec.teacher("שרה לוי").latest_period_on, {4: 4})

    def test_a_split_keeps_its_order_and_its_repeats(self):
        spec = _demo()
        maths = [l for l in spec.lessons if l.subject == "חשבון"]
        self.assertTrue(maths)
        for lesson in maths:
            self.assertEqual(lesson.pattern, [2, 1, 1, 1])

    def test_a_school_survives_export_and_reimport(self):
        spec = _demo()
        again, report = importer.read_files(samples.spec_to_csv(spec))
        self.assertTrue(report.ok, report.errors)
        self.assertEqual(spec_mod.to_dict(again), spec_mod.to_dict(spec))

    def _read(self, text: str, extra: str = ""):
        files = [("data.csv", text.encode("utf-8"))]
        if extra:
            files.append(("more.csv", extra.encode("utf-8")))
        return importer.read_files(files)

    def test_an_unknown_day_is_reported_with_its_row_number(self):
        spec, report = self._read(
            "## מורות\nשם המורה\nרות\n"
            "## כיתות\nשם הכיתה,יום,מינימום שעות,מקסימום שעות\nא,כל הימים,4,5\n"
            "## שיעורים\nכיתה,מקצוע,מורות,שעות\nא,תורה,רות,4\n"
            "## אילוצים\nשם המורה,סוג,יום,שעות\nרות,יום חופש,יום ד,\n")
        self.assertIsNone(spec)
        self.assertTrue(any("שורה 2" in e and "יום ד" in e
                            for e in report.errors), report.errors)

    def test_a_lesson_naming_someone_off_the_staff_list_is_reported(self):
        spec, report = self._read(
            "## מורות\nשם המורה\nרות\n"
            "## כיתות\nשם הכיתה,יום,מינימום שעות,מקסימום שעות\nא,כל הימים,4,5\n"
            "## שיעורים\nכיתה,מקצוע,מורות,שעות\nא,תורה,מישהי אחרת,4\n")
        self.assertIsNone(spec)
        self.assertTrue(any("מישהי אחרת" in e for e in report.errors),
                        report.errors)

    def test_an_unreadable_table_is_named_rather_than_dropped(self):
        _, report = self._read("## משהו אחר\nעמודה,עמודה\n1,2\n")
        self.assertTrue(any("לא זוהה סוג הטבלה" in w
                            for w in report.warnings), report.warnings)

    def test_a_full_backup_comes_back_with_its_timetable(self):
        spec = _demo()
        payload = {"school": spec_mod.to_dict(spec),
                   "timetable": tt_mod.Document.blank(spec).to_dict(),
                   "settings": {"overrides": {"teacher_spread":
                                              {"weight": 0}}}}
        again, report = importer.read_files(
            [("backup.json", json.dumps(payload).encode("utf-8"))])
        self.assertTrue(report.ok, report.errors)
        self.assertEqual(spec_mod.to_dict(again), spec_mod.to_dict(spec))
        self.assertIn("timetable", report.restored)


class TestXlsx(unittest.TestCase):
    def _workbook(self, sheets: dict[str, list[list[str]]]) -> bytes:
        """A minimal .xlsx, written by hand, to read back."""
        strings: list[str] = []

        def index(text: str) -> int:
            if text not in strings:
                strings.append(text)
            return strings.index(text)

        parts = {}
        for n, (name, rows) in enumerate(sheets.items(), start=1):
            body = ""
            for r, row in enumerate(rows, start=1):
                cells = "".join(
                    f'<c r="{chr(64 + c)}{r}" t="s"><v>{index(v)}</v></c>'
                    for c, v in enumerate(row, start=1))
                body += f'<row r="{r}">{cells}</row>'
            parts[f"xl/worksheets/sheet{n}.xml"] = (
                '<?xml version="1.0"?><worksheet xmlns="http://schemas.'
                'openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
                f'{body}</sheetData></worksheet>')
        ns = ('xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/'
              'main" xmlns:r="http://schemas.openxmlformats.org/'
              'officeDocument/2006/relationships"')
        tabs = "".join(f'<sheet name="{name}" sheetId="{n}" r:id="rId{n}"/>'
                       for n, name in enumerate(sheets, start=1))
        buf = pathlib.Path(tempfile.mkdtemp()) / "book.xlsx"
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("xl/workbook.xml",
                        f'<?xml version="1.0"?><workbook {ns}>'
                        f'<sheets>{tabs}</sheets></workbook>')
            zf.writestr("xl/_rels/workbook.xml.rels",
                        '<?xml version="1.0"?><Relationships xmlns="http://'
                        'schemas.openxmlformats.org/package/2006/'
                        'relationships">' + "".join(
                            f'<Relationship Id="rId{n}" Target="worksheets/'
                            f'sheet{n}.xml"/>'
                            for n in range(1, len(sheets) + 1))
                        + "</Relationships>")
            zf.writestr("xl/sharedStrings.xml",
                        '<?xml version="1.0"?><sst xmlns="http://schemas.'
                        'openxmlformats.org/spreadsheetml/2006/main">'
                        + "".join(f"<si><t>{s}</t></si>" for s in strings)
                        + "</sst>")
            for path, text in parts.items():
                zf.writestr(path, text)
        return buf.read_bytes()

    def test_a_workbook_with_a_sheet_per_table_imports(self):
        data = self._workbook({
            "מורות": [["שם המורה", "מחנכת של"], ["רות", "א"]],
            "כיתות": [["שם הכיתה", "יום", "מינימום שעות", "מקסימום שעות"],
                      ["א", "כל הימים", "3", "5"]],
            "שיעורים": [["כיתה", "מקצוע", "מורות", "שעות"],
                        ["א", "תורה", "רות", "4"]],
        })
        spec, report = importer.read_files([("book.xlsx", data)])
        self.assertTrue(report.ok, report.errors)
        self.assertEqual(spec.teacher_names, ["רות"])
        self.assertEqual(spec.lessons[0].hours, 4)

    def test_a_file_that_is_not_a_workbook_says_so_in_hebrew(self):
        _, report = importer.read_files([("book.xlsx", b"not a zip")])
        self.assertTrue(any("Excel" in e for e in report.errors),
                        report.errors)


class TestWordImport(unittest.TestCase):
    """A .docx is read with the standard library, so this works whether or
    not python-docx is installed — but writing the fixture needs it."""

    def _document(self, tables: list[tuple[str, list[list]]]) -> bytes:
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx is not installed")
        import io
        doc = Document()
        for title, rows in tables:
            doc.add_heading(title, level=1)
            table = doc.add_table(rows=0, cols=max(len(r) for r in rows))
            for row in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = "" if value is None else str(value)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_a_heading_names_the_table_under_it(self):
        from app import docx_read
        data = self._document([("מורות", [["שם המורה"], ["רות"]]),
                               ("כיתות", [["שם הכיתה"], ["א"]])])
        got = docx_read.read_tables(data)
        self.assertEqual(list(got), ["מורות", "כיתות"])
        self.assertEqual(got["מורות"], [["שם המורה"], ["רות"]])

    def test_a_word_school_imports(self):
        data = self._document([
            ("מורות", [["שם המורה", "מחנכת של"], ["רות", "א"]]),
            ("כיתות", [["שם הכיתה", "יום", "מינימום שעות", "מקסימום שעות"],
                       ["א", "כל הימים", 3, 5]]),
            ("שיעורים", [["כיתה", "מקצוע", "מורות", "שעות"],
                         ["א", "תורה", "רות", 4]]),
            ("אילוצים", [["שם המורה", "סוג", "יום", "שעות"],
                         ["רות", "יום חופש", "שני", ""]]),
        ])
        spec, report = importer.read_files([("school.docx", data)])
        self.assertTrue(report.ok, report.errors)
        self.assertEqual(spec.teacher("רות").off_fixed, [1])
        self.assertEqual(spec.lessons[0].hours, 4)

    def test_the_word_example_file_reimports_cleanly(self):
        data = samples.template_docx()
        if data is None:
            self.skipTest("python-docx is not installed")
        spec, report = importer.read_files([(samples.WORD_FILE, data)])
        self.assertTrue(report.ok, report.errors)
        self.assertEqual(report.warnings, [])
        # ...and describes the same school as the CSV example.
        csv_spec, _ = samples.demo_spec()
        spec.name, spec.year = csv_spec.name, csv_spec.year
        self.assertEqual(spec_mod.to_dict(spec), spec_mod.to_dict(csv_spec))

    def test_an_old_binary_doc_says_what_to_do(self):
        _, report = importer.read_files([("staff.doc", b"\xd0\xcf\x11\xe0")])
        self.assertTrue(any("docx" in e for e in report.errors),
                        report.errors)

    def test_something_that_is_not_a_document_is_reported(self):
        _, report = importer.read_files([("staff.docx", b"nonsense")])
        self.assertTrue(any("Word" in e for e in report.errors),
                        report.errors)


class TestConstraintScreen(unittest.TestCase):
    def test_every_soft_entry_names_labels_the_model_really_prices(self):
        """A catalogue entry whose labels the solver never emits would show a
        priority control that changes nothing."""
        spec = activate.capture()
        with activate.activated(spec):
            built = solver.Scheduler(spec.requirements(), time_limit=1)
            priced = set(built._penalty_terms)
            listed = {label for c in cat.CATALOGUE if c.kind == "soft"
                      for label in c.labels(spec)}
        # Labels the model emits only in a configuration this school is not
        # in: the anchor rule needs an anchor, the two window labels are
        # alternatives to the ones it does use, §11.2(ג) is switched off, and
        # the extra-opening nudge appears only where preferred > minimum.
        conditional = {"שינוי מהשיבוץ שאושר", "חלונות", "חלונות מעבר לאחד",
                       "יום חמישי: חלון ביום זה",
                       "מחנכת: פתיחה מועדפת נוספת"}
        self.assertEqual((listed - priced) - conditional, set())

    def test_disabling_a_hard_rule_produces_its_relax_key(self):
        settings = cat.Settings()
        self.assertEqual(settings.relax(), frozenset())
        settings.set("contiguity", enabled=False)
        self.assertIn("contiguity", settings.relax())
        self.assertIn("contiguity", settings.relax_ids())

    def test_a_locked_rule_cannot_be_switched_off(self):
        with self.assertRaises(ValueError):
            cat.Settings().set("single_booking", enabled=False)

    def test_zeroing_a_soft_rule_removes_it_from_the_objective(self):
        spec = activate.capture()
        settings = cat.Settings()
        settings.set("teacher_spread", enabled=False)
        label = "פיזור: 2 ימים קצרים למורה"
        self.assertEqual(settings.weights(spec)[label], 0)
        with activate.activated(spec):
            built = solver.Scheduler(spec.requirements(), time_limit=1,
                                     weights=settings.weights(spec))
        self.assertNotIn(label, built._penalty_terms)

    def test_raising_a_priority_reaches_the_model(self):
        spec = activate.capture()
        settings = cat.Settings()
        settings.set("teacher_spread", weight=solver.W_STRONG)
        with activate.activated(spec):
            built = solver.Scheduler(spec.requirements(), time_limit=1,
                                     weights=settings.weights(spec))
        weights = {w for _, w in
                   built._penalty_terms["פיזור: 2 ימים קצרים למורה"]}
        self.assertEqual(weights, {solver.W_STRONG})

    def test_every_field_writes_back_what_it_reads(self):
        """A field whose `put` does not invert its `get` would change the
        school the moment someone saves a form they did not touch."""
        spec = activate.capture()
        before = spec_mod.to_dict(spec)
        for c in cat.CATALOGUE:
            if not c.fields:
                continue
            values = {f.key: f.get(spec) for f in c.fields}
            with self.subTest(constraint=c.id):
                self.assertEqual(cat.apply_fields(spec, c.id, values), [])
                self.assertEqual(spec_mod.to_dict(spec), before)

    def test_an_edited_name_list_reaches_the_school(self):
        spec = activate.capture()
        name = spec.teacher_names[0]
        cat.apply_fields(spec, "no_window_teachers",
                         {"no_window_teachers": [name]})
        self.assertEqual(spec.rules.no_window_teachers, [name])

    def test_a_name_the_school_does_not_have_is_refused(self):
        spec = activate.capture()
        with self.assertRaises(ValueError):
            cat.apply_fields(spec, "no_window_teachers",
                             {"no_window_teachers": ["אין מורה כזו"]})
        with self.assertRaises(ValueError):
            cat.apply_fields(spec, "no_window_teachers", {"no_such_field": []})

    def test_a_refused_edit_leaves_the_saved_school_untouched(self):
        from app import server
        api = server.Workspaces(tempfile.mkdtemp(), auth=False).api(None)
        pid = api.post_demo({})["id"]
        before = spec_mod.to_dict(api.project(pid).spec)
        with self.assertRaises(server.HttpError) as caught:
            api.post_constraints({"id": pid, "constraint": "no_window_teachers",
                                  "params": {"no_window_teachers":
                                             ["אין מורה כזו"]}})
        self.assertEqual(caught.exception.status, 400)
        self.assertEqual(spec_mod.to_dict(api.project(pid).spec), before)

        name = api.project(pid).spec.teacher_names[0]
        api.post_constraints({"id": pid, "constraint": "no_window_teachers",
                              "params": {"no_window_teachers": [name]}})
        self.assertEqual(api.project(pid).spec.rules.no_window_teachers,
                         [name])

    def test_the_listing_covers_every_category(self):
        spec = activate.capture()
        groups = cat.listing(spec, cat.Settings())
        self.assertEqual([g["id"] for g in groups], cat.CATEGORY_ORDER)
        self.assertTrue(all(g["constraints"] for g in groups))


class TestDocument(unittest.TestCase):
    def setUp(self):
        self.spec = _demo()
        self.doc = tt_mod.Document.blank(self.spec)

    def test_row_ids_survive_a_deletion(self):
        keep = self.doc.rows[5].uid
        self.doc.remove(self.doc.rows[0].uid)
        self.assertEqual(self.doc.row(keep).uid, keep)

    def test_coverage_notices_a_deleted_lesson(self):
        self.doc.remove(self.doc.rows[0].uid)
        gaps = tt_mod.coverage(self.spec, self.doc)
        self.assertEqual(len(gaps), 1)
        self.assertEqual(gaps[0]["present"], gaps[0]["required"] - 1)

    def test_coverage_is_silent_when_nothing_moved(self):
        self.assertEqual(tt_mod.coverage(self.spec, self.doc), [])

    def test_it_survives_a_trip_through_json(self):
        self.doc.move(self.doc.rows[0].uid, 0, 1)
        self.doc.mark_approved()
        again = tt_mod.Document.from_dict(json.loads(json.dumps(
            self.doc.to_dict())))
        self.assertEqual(again.to_dict(), self.doc.to_dict())


class TestSolveAndEdit(unittest.TestCase):
    """One end-to-end pass over the small school: build it, then edit it."""

    @classmethod
    def setUpClass(cls):
        cls.spec = _demo()
        cls.settings = cat.Settings()
        cls.doc, cls.report = engine.generate(
            cls.spec, cls.settings, time_limit=90, workers=4)

    def test_it_solves_and_the_independent_validator_agrees(self):
        self.assertIsNotNone(self.doc, self.report.attempts)
        self.assertTrue(self.report.hard_ok, self.report.problems)
        self.assertEqual(self.report.errors, 0)

    def test_every_lesson_is_placed(self):
        self.assertTrue(all(r.placed for r in self.doc.rows))
        self.assertEqual(tt_mod.coverage(self.spec, self.doc), [])

    def test_alternatives_never_offer_a_slot_that_breaks_a_hard_rule(self):
        """The claim the screen makes about a green option, checked."""
        with activate.activated(self.spec):
            for row in self.doc.rows[:8]:
                found = alt.find(self.spec, self.settings, self.doc, row.uid)
                for option in found["options"]:
                    trial = self.doc.copy()
                    trial.exchange(row.uid, option["day"], option["period"],
                                   option["partner"])
                    got = tt_mod.review(self.spec, self.settings, trial,
                                        require_complete=False)
                    self.assertEqual(got["violations"], [],
                                     f"{row.subject} → {option}")
                    self.assertEqual(got["total"] - found["base_total"],
                                     option["delta"])

    def test_a_move_onto_an_occupied_slot_is_reported_as_occupied(self):
        with activate.activated(self.spec):
            row = self.doc.rows[0]
            found = alt.find(self.spec, self.settings, self.doc, row.uid)
        occupied = [o for o in found["blocked"]
                    if any(b["rule"] == "single_booking" for b in o["blocking"])]
        self.assertTrue(occupied)

    def test_a_bad_manual_move_names_the_rule_it_broke(self):
        with activate.activated(self.spec):
            after = self.doc.copy()
            row = next(r for r in after.rows if r.period == 1)
            after.move(row.uid, row.day, 6)
            outcome = alt.check_edit(self.spec, self.settings, self.doc, after)
        self.assertFalse(outcome["ok"])
        rules = {v["rule"] for v in outcome["introduced"]}
        self.assertTrue({"contiguity", "class_span"} & rules, rules)
        self.assertTrue(all(v["title"] for v in outcome["introduced"]))

    def test_removing_a_lesson_leaves_the_rest_legal(self):
        with activate.activated(self.spec):
            after = self.doc.copy()
            after.unplace(after.rows[0].uid)
            got = tt_mod.review(self.spec, self.settings, after,
                                require_complete=False)
        self.assertEqual(len(got["unplaced"]), 1)
        self.assertIn("uid", got["unplaced"][0])


class TestStore(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.store = store_mod.Store(self.dir)

    def test_a_project_round_trips(self):
        spec = _demo()
        project = self.store.create(spec)
        project.doc = tt_mod.Document.blank(spec)
        project.settings.set("teacher_spread", weight=0)
        self.store.save(project)
        again = self.store.load(project.id)
        self.assertEqual(spec_mod.to_dict(again.spec), spec_mod.to_dict(spec))
        self.assertEqual(again.settings.to_dict(),
                         project.settings.to_dict())
        self.assertEqual(len(again.doc.rows), len(project.doc.rows))

    def test_undo_returns_the_timetable_that_was_showing(self):
        spec = _demo()
        project = self.store.create(spec)
        project.doc = tt_mod.Document.blank(spec)
        project.doc.move(project.doc.rows[0].uid, 0, 1)
        self.store.save(project)
        before = project.doc.to_dict()

        self.store.snapshot(project, "לפני עריכה")
        project.doc.move(project.doc.rows[0].uid, 1, 3)
        self.store.save(project)

        self.store.undo(project)
        self.assertEqual(project.doc.to_dict(), before)
        with self.assertRaises(KeyError):
            self.store.undo(project)

    def test_a_stray_separator_cannot_escape_the_store(self):
        with self.assertRaises(KeyError):
            self.store.load("../../etc")


class TestExport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.spec = activate.capture()
        saved = json.loads((ROOT / "solution.json").read_text(encoding="utf-8"))
        with activate.activated(cls.spec):
            cls.tt = Timetable({int(k): tuple(v)
                                for k, v in saved["placement"].items()},
                               expand(cls.spec.requirements()))

    def _write(self, **kw) -> bytes:
        try:
            import word
        except ImportError:
            self.skipTest("python-docx is not installed")
        path = pathlib.Path(tempfile.mkdtemp()) / "sheet.docx"
        with activate.activated(self.spec):
            word.sheet_docx(self.tt, "מערכת שעות", path, **kw)
        return path.read_bytes()

    def test_the_final_version_is_marked_read_only(self):
        with zipfile.ZipFile(pathlib.Path(tempfile.mkdtemp()) / "x", "w"):
            pass
        data = self._write(locked=True)
        with zipfile.ZipFile(__import__("io").BytesIO(data)) as zf:
            settings = zf.read("word/settings.xml").decode()
        self.assertIn('w:edit="readOnly"', settings)
        self.assertIn('w:enforcement="1"', settings)

    def test_the_editable_version_is_not(self):
        data = self._write()
        with zipfile.ZipFile(__import__("io").BytesIO(data)) as zf:
            settings = zf.read("word/settings.xml").decode()
        self.assertNotIn("documentProtection", settings)

    def test_one_view_is_smaller_than_both(self):
        self.assertLess(len(self._write(view="teacher")),
                        len(self._write(view="both")))


class TestServerRoutes(unittest.TestCase):
    """The API, exercised through `Api` rather than a socket."""

    def setUp(self):
        from app import server
        self.workspaces = server.Workspaces(tempfile.mkdtemp(), auth=False)
        self.api = self.workspaces.api(None)
        self.server = server

    def test_the_demo_school_can_be_created_and_read(self):
        created = self.api.post_demo({})
        state = self.api.get_school({"id": created["id"]})
        self.assertEqual(len(state["teachers"]), 4)
        self.assertEqual(state["grid"]["days"][0], "ראשון")
        self.assertNotIn("rows", state)

    def test_uploading_the_example_files_creates_a_school(self):
        out = self.api.post_import({"files": samples.template_files(),
                                    "name": "בדיקה", "year": ""})
        self.assertTrue(out["ok"], out.get("report"))
        self.assertEqual(out["checks"]["errors"], 0)

    def test_an_unknown_school_is_a_404_not_a_crash(self):
        with self.assertRaises(self.server.HttpError) as caught:
            self.api.get_school({"id": "nope"})
        self.assertEqual(caught.exception.status, 404)

    def test_asking_to_export_before_solving_says_so(self):
        created = self.api.post_demo({})
        project = self.api.project(created["id"])
        with self.assertRaises(self.server.HttpError):
            self.server.export(project, "html")

    def test_the_template_download_is_a_readable_zip(self):
        created = self.api.post_demo({})
        name, ctype, data = self.server.export(
            self.api.project(created["id"]), "template")
        self.assertEqual(ctype, "application/zip")
        with zipfile.ZipFile(__import__("io").BytesIO(data)) as zf:
            names = zf.namelist()
        # The two CSVs always; the same example as a Word document only where
        # python-docx is installed to write one.
        self.assertEqual(names[:2], [samples.TEACHERS_FILE,
                                     samples.LIMITS_FILE])
        self.assertEqual(len(names), 3 if samples.template_docx() else 2)


class TestAccounts(unittest.TestCase):
    def setUp(self):
        from app import users as users_mod
        self.mod = users_mod
        self.users = users_mod.Users(tempfile.mkdtemp())

    def _make(self, name="rivka", password="a-good-password"):
        return self.users.register(name, password, "רבקה")

    def test_a_password_is_never_stored(self):
        user = self._make()
        blob = json.dumps(user.to_dict(), ensure_ascii=False)
        self.assertNotIn("a-good-password", blob)
        self.assertNotIn("a-good-password",
                         (self.users.root / "users.json")
                         .read_text(encoding="utf-8"))

    def test_the_right_password_authenticates_and_a_wrong_one_does_not(self):
        self._make()
        self.assertEqual(self.users.authenticate("rivka",
                                                 "a-good-password").username,
                         "rivka")
        with self.assertRaises(ValueError):
            self.users.authenticate("rivka", "wrong-password-here")

    def test_an_unknown_name_and_a_wrong_password_say_the_same_thing(self):
        """The login screen must not be a way to find out who has an account."""
        self._make()
        try:
            self.users.authenticate("nobody", "whatever-it-is")
        except ValueError as exc:
            unknown = str(exc)
        try:
            self.users.authenticate("rivka", "whatever-it-is")
        except ValueError as exc:
            wrong = str(exc)
        self.assertEqual(unknown, wrong)

    def test_the_same_name_cannot_be_registered_twice(self):
        self._make("Rivka")
        with self.assertRaises(ValueError):
            self._make("rivka")          # case- and NFKC-folded to one key

    def test_a_weak_account_is_refused_with_a_reason(self):
        for name, password in (("ab", "a-good-password"),
                               ("rivka", "short"),
                               ("riv ka", "a-good-password")):
            with self.assertRaises(ValueError):
                self.users.register(name, password)

    def test_repeated_failures_lock_the_account(self):
        self._make()
        for _ in range(self.mod.MAX_ATTEMPTS):
            with self.assertRaises(ValueError):
                self.users.authenticate("rivka", "not-the-password")
        with self.assertRaises(ValueError) as caught:
            self.users.authenticate("rivka", "a-good-password")
        self.assertIn("נעול", str(caught.exception))

    def test_a_session_round_trips_and_can_be_ended(self):
        user = self._make()
        token = self.users.start_session(user)
        self.assertEqual(self.users.session_user(token).id, user.id)
        self.users.end_session(token)
        self.assertIsNone(self.users.session_user(token))
        self.assertIsNone(self.users.session_user("made-up"))

    def test_changing_a_password_ends_every_other_session(self):
        user = self._make()
        elsewhere = self.users.start_session(user)
        self.users.change_password(user, "a-good-password", "another-good-one")
        self.assertIsNone(self.users.session_user(elsewhere))
        self.assertEqual(self.users.authenticate("rivka",
                                                 "another-good-one").id, user.id)

    def test_the_cookie_is_secure_only_behind_a_tls_proxy(self):
        """Always-`Secure` would lock out plain http://127.0.0.1; never-
        `Secure` would let a deployment behind HTTPS leak the session."""
        import email.message
        from app import server

        def cookies(proto):
            handler = server.Handler.__new__(server.Handler)
            handler.headers = email.message.Message()
            if proto:
                handler.headers["X-Forwarded-Proto"] = proto
            handler._cookies = []
            handler._set_cookie("token")
            handler._set_cookie(None)
            return handler._cookies

        for cookie in cookies("https"):
            self.assertIn("; Secure", cookie)
        for proto in ("", "http"):
            for cookie in cookies(proto):
                self.assertNotIn("Secure", cookie)

    def test_two_accounts_never_see_each_other_schools(self):
        from app import server
        root = tempfile.mkdtemp()
        workspaces = server.Workspaces(root, auth=True)
        one = workspaces.users.register("aviva", "a-good-password")
        two = workspaces.users.register("bracha", "a-good-password")
        workspaces.api(one).post_demo({})
        self.assertEqual(len(workspaces.api(one).get_schools({})["schools"]), 1)
        self.assertEqual(len(workspaces.api(two).get_schools({})["schools"]), 0)


class TestMultipart(unittest.TestCase):
    def test_a_file_arrives_byte_for_byte(self):
        from app.server import _parse_multipart
        payload = "שלום,עולם\r\n".encode("utf-8-sig")
        body = (b"--X\r\nContent-Disposition: form-data; name=\"name\"\r\n\r\n"
                b"bet\r\n"
                b"--X\r\nContent-Disposition: form-data; name=\"files\"; "
                b"filename=\"a.csv\"\r\nContent-Type: text/csv\r\n\r\n"
                + payload + b"\r\n--X--\r\n")
        got = _parse_multipart(body, b"X")
        self.assertEqual(got["name"], ["bet"])
        self.assertEqual(got["files"], [("a.csv", payload)])


if __name__ == "__main__":
    unittest.main()
