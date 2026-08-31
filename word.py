# -*- coding: utf-8 -*-
"""The hand-out as a Word document.

Same content as `report.sheet_page` — the ten class grids, then the
seventeen teacher grids — but as a real .docx: Word tables the school can
edit, not an HTML page renamed.  Everything here is layout; the timetable
itself is read through the same `report._cell` the HTML uses, so the two
hand-outs cannot disagree about what is scheduled.

Hebrew needs three separate right-to-left switches, and missing any one of
them reads as a bug in the timetable rather than in the formatting: the
section (page direction), the table (`bidiVisual`, which mirrors the column
order) and every paragraph and run inside the cells.

The grids are built as XML text and parsed once per table rather than
assembled element by element.  All twenty-seven tables have the same shape,
and driving ~1500 cells through the python-docx object model costs about a
minute and a half per document — `table.cell(i, j)` rebuilds the table's
whole cell map on every call — where one parse per table costs about a
second.  The thirty headings use the ordinary API, where it does not matter.
"""
from __future__ import annotations

from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt

import report
import school
from model import DAYS, PERIODS_PER_DAY, Timetable

#: Kept in step with report._CSS so the two hand-outs look like one document.
HEAD_FILL = "EFECE5"      # day header and period column
LESSON_FILL = "FFFFFF"    # a period that has a lesson in it
FREE_FILL = "F5F3EF"      # a period the class or teacher does not use
GAP_FILL = "FBE9E7"       # a teacher's window — the thing to look at twice
OUTSIDE_FILL = "E4E0D8"   # a period that does not exist that day (Friday)
MUTED = "6B6862"

FONT = "Arial"

#: Column widths in twentieths of a point: the period number, then a day.
_W_PERIOD, _W_DAY = 620, 2380


def _run(text: str, *, size: float, bold=False, colour=None,
         break_first=False) -> str:
    """One right-to-left run.  `w:sz` is in half-points."""
    props = [f'<w:rFonts w:ascii="{FONT}" w:hAnsi="{FONT}" w:cs="{FONT}"/>',
             "<w:rtl/>"]
    if bold:
        props.append("<w:b/><w:bCs/>")
    if colour:
        props.append(f'<w:color w:val="{colour}"/>')
    half = int(round(size * 2))
    props.append(f'<w:sz w:val="{half}"/><w:szCs w:val="{half}"/>')
    return (f'<w:r><w:rPr>{"".join(props)}</w:rPr>'
            f'{"<w:br/>" if break_first else ""}'
            f'<w:t xml:space="preserve">{escape(text)}</w:t></w:r>')


def _tc(runs: str, fill: str, width: int = _W_DAY) -> str:
    """One cell: shaded, vertically centred, holding a single paragraph."""
    return (f'<w:tc><w:tcPr>'
            f'<w:tcW w:w="{width}" w:type="dxa"/>'
            f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>'
            f'<w:vAlign w:val="center"/></w:tcPr>'
            f'<w:p><w:pPr><w:bidi/><w:jc w:val="center"/>'
            f'<w:spacing w:before="20" w:after="20" w:line="240" '
            f'w:lineRule="auto"/></w:pPr>{runs}</w:p></w:tc>')


def _grid_xml(grid, gaps, mode) -> str:
    periods = max(PERIODS_PER_DAY)
    cols = "".join(f'<w:gridCol w:w="{_W_PERIOD if i == 0 else _W_DAY}"/>'
                   for i in range(len(DAYS) + 1))

    head = [_tc(_run("שעה", size=9, bold=True), HEAD_FILL, _W_PERIOD)]
    head += [_tc(_run(name, size=10, bold=True), HEAD_FILL) for name in DAYS]
    # tblHeader repeats the day names if a grid falls across a page break.
    rows = [f'<w:tr><w:trPr><w:tblHeader/></w:trPr>{"".join(head)}</w:tr>']

    for p in range(1, periods + 1):
        cells = [_tc(_run(str(p), size=9, bold=True), HEAD_FILL, _W_PERIOD)]
        for d in range(len(DAYS)):
            if p > PERIODS_PER_DAY[d]:
                cells.append(_tc("", OUTSIDE_FILL))
                continue
            subject, who = report._cell(grid.get((d, p)), mode)
            if not subject:
                cells.append(_tc(_run("—", size=9, colour=MUTED),
                                 GAP_FILL if (d, p) in gaps else FREE_FILL))
                continue
            # The second line names what the heading does not already say:
            # the teacher in a class grid, the class in a teacher grid.
            cells.append(_tc(_run(subject, size=9)
                             + _run(who, size=7.5, colour=MUTED,
                                    break_first=True),
                             LESSON_FILL))
        rows.append(f'<w:tr>{"".join(cells)}</w:tr>')

    return (f'<w:tbl {nsdecls("w")}><w:tblPr>'
            f'<w:tblStyle w:val="TableGrid"/>'
            f'<w:tblW w:w="0" w:type="auto"/>'
            f'<w:jc w:val="center"/><w:bidiVisual/>'
            f'</w:tblPr><w:tblGrid>{cols}</w:tblGrid>'
            f'{"".join(rows)}</w:tbl>')


def _heading(doc, text: str, level: int):
    par = doc.add_heading(level=level)
    run = par.add_run(text)
    run.font.name = FONT
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    par._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    rPr = run._element.get_or_add_rPr()
    rPr.append(OxmlElement("w:rtl"))
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rPr.insert(0, fonts)
    fonts.set(qn("w:cs"), FONT)
    return par


def _grid(doc, grid, gaps, mode):
    """Append one grid where the reader expects it — under its heading.

    Not `body.append`: the body ends with `w:sectPr`, and python-docx inserts
    every paragraph *before* it while a bare append lands *after* it.  Mixing
    the two puts every heading first and every table last, which is a
    document nobody can read.
    """
    body = doc.element.body
    tbl = parse_xml(_grid_xml(grid, gaps, mode))
    sectPr = body.find(qn("w:sectPr"))
    if sectPr is None:
        body.append(tbl)
    else:
        sectPr.addprevious(tbl)
    doc.add_paragraph()


def _protect(doc) -> None:
    """Mark the document read-only in Word.

    Word's own "editing restricted / no changes" flag.  It preserves the
    layout for the person the timetable is handed to, and it is honest about
    what it is: no password, so a reader who means to edit can turn it off in
    Word — which is the right level for a hand-out, and is what the export
    screen says.
    """
    settings = doc.settings.element
    node = parse_xml(
        '<w:documentProtection xmlns:w="http://schemas.openxmlformats.org/'
        'wordprocessingml/2006/main" w:edit="readOnly" w:enforcement="1"/>')
    # `w:documentProtection` sits before `w:defaultTabStop` in the schema's
    # sequence, and Word rejects a settings part whose children are out of
    # order — so it is inserted, not appended.
    anchor = settings.find(qn("w:defaultTabStop"))
    if anchor is None:
        settings.append(node)
    else:
        anchor.addprevious(node)


def sheet_docx(tt: Timetable, title: str, path, locked: bool = False,
               view: str = "both"):
    """Write the hand-out to `path`.

    Landscape, because seven columns of Hebrew subject names do not fit
    across a portrait page without wrapping every cell.  `view` selects the
    class grids, the teacher grids or both; `locked` produces the version
    meant for distribution rather than for further editing.
    """
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = (section.page_height,
                                               section.page_width)
    for side in ("top", "bottom", "left", "right"):
        setattr(section, f"{side}_margin", Cm(1.2))
    section._sectPr.append(OxmlElement("w:bidi"))

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(9)

    _heading(doc, title, 0)

    if view in ("both", "class"):
        _heading(doc, "מערכות הכיתות", 1)
        for c in school.CLASSES:
            _heading(doc, f"כיתה {c.name}", 2)
            _grid(doc, tt.by_class(c.name), set(), "class")

    if view == "both":
        doc.add_page_break()
    if view in ("both", "teacher"):
        _heading(doc, "מערכות המורות", 1)
    for t in (school.TEACHERS if view in ("both", "teacher") else []):
        grid = tt.by_teacher(t.name)
        gaps = set()
        for d in range(len(DAYS)):
            busy = sorted(p for (dd, p) in grid if dd == d)
            if busy:
                gaps |= {(d, p) for p in range(busy[0], busy[-1] + 1)
                         if (d, p) not in grid}
        # A homeroom teacher's two חינוך hours occupy no period, so they are
        # named beside the count rather than folded into it.
        extra = (f" + {school.HOMEROOM_EDUCATION_HOURS} חינוך"
                 if t.homeroom else "")
        _heading(doc, f"{t.name} — {len(grid)} שעות{extra}", 2)
        _grid(doc, grid, gaps, "teacher")

    if locked:
        _protect(doc)
    doc.save(str(path))
    return path
